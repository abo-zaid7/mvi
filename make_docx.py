"""
Convert REPORT.md into a Microsoft Word document.

Run with:   ./mvi-env/bin/python make_docx.py

Handles the subset of Markdown the report actually uses: headings, pipe tables,
fenced code blocks, bullet lists, block quotes, horizontal rules, and inline
**bold** / *italic* / `code`.  Tables come out as real Word tables so they can be
restyled in Word afterwards.
"""

import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Inches

SOURCE = "REPORT.md"
TARGET = "Medical Image Segmentation Report.docx"


# ------------------------------------------------------------------ inline text
def add_runs(paragraph, text):
    """Split a line into **bold**, *italic* and `code` runs."""
    pattern = r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)"
    for piece in re.split(pattern, text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            paragraph.add_run(piece[2:-2]).bold = True
        elif piece.startswith("*") and piece.endswith("*"):
            paragraph.add_run(piece[1:-1]).italic = True
        elif piece.startswith("`") and piece.endswith("`"):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0xC0, 0x30, 0x30)
        else:
            paragraph.add_run(piece)


def clean(cell):
    """Strip markdown emphasis from table cells, keeping a bold flag."""
    text = cell.strip()
    bold = "**" in text
    text = text.replace("**", "").replace("`", "")
    text = re.sub(r"(?<!\*)\*(?!\*)", "", text)
    return text, bold


# ------------------------------------------------------------------ blocks
def add_table(document, rows):
    header = [c.strip() for c in rows[0].strip().strip("|").split("|")]
    body = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows[2:]]

    table = document.add_table(rows=1, cols=len(header))
    table.style = "Light Grid Accent 1"
    table.autofit = True

    for i, name in enumerate(header):
        text, _ = clean(name)
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.size = Pt(9)

    for line in body:
        cells = table.add_row().cells
        for i, value in enumerate(line[:len(header)]):
            text, bold = clean(value)
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(text)
            run.bold = bold
            run.font.size = Pt(9)

    document.add_paragraph()


def add_code(document, lines):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.3)
    paragraph.paragraph_format.space_after = Pt(10)
    run = paragraph.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)


def convert(source=SOURCE, target=TARGET):
    with open(source) as handle:
        lines = handle.read().split("\n")

    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        # ---- fenced code
        if stripped.startswith("```"):
            index += 1
            block = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index])
                index += 1
            add_code(document, block)
            index += 1
            continue

        # ---- tables
        if stripped.startswith("|") and index + 1 < len(lines) and \
                re.match(r"^\|[\s:\-|]+\|$", lines[index + 1].strip()):
            block = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                block.append(lines[index])
                index += 1
            add_table(document, block)
            continue

        # ---- headings
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip()
            if level == 1:
                heading = document.add_heading(text.replace("**", ""), level=0)
                heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                document.add_heading(text.replace("**", ""), level=min(level - 1, 4))
            index += 1
            continue

        # ---- horizontal rule
        if stripped in ("---", "***", "___"):
            document.add_paragraph("_" * 68).alignment = WD_ALIGN_PARAGRAPH.CENTER
            index += 1
            continue

        # ---- block quote
        if stripped.startswith(">"):
            block = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                block.append(lines[index].strip().lstrip(">").strip())
                index += 1
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.4)
            add_runs(paragraph, " ".join(block))
            for run in paragraph.runs:
                run.italic = True
            continue

        # ---- bullets and numbered items
        bullet = re.match(r"^(\s*)[-*]\s+(.*)", line)
        number = re.match(r"^(\s*)(\d+)\.\s+(.*)", line)
        if bullet or number:
            indent = len(bullet.group(1) if bullet else number.group(1))
            text = bullet.group(2) if bullet else number.group(3)
            style_name = "List Bullet" if bullet else "List Number"
            if indent >= 2:
                style_name += " 2"
            paragraph = document.add_paragraph(style=style_name)
            add_runs(paragraph, text)
            index += 1
            # keep wrapped continuation lines with their bullet
            while index < len(lines) and lines[index].startswith("  ") and \
                    lines[index].strip() and not re.match(r"^\s*[-*\d]", lines[index]) and \
                    not lines[index].strip().startswith("|"):
                paragraph.add_run(" ")
                add_runs(paragraph, lines[index].strip())
                index += 1
            continue

        # ---- blank
        if not stripped:
            index += 1
            continue

        # ---- ordinary paragraph, joining wrapped lines
        block = [stripped]
        index += 1
        while index < len(lines) and lines[index].strip() and \
                not re.match(r"^(#|\||>|```|---|\s*[-*]\s|\s*\d+\.\s)", lines[index]):
            block.append(lines[index].strip())
            index += 1
        add_runs(document.add_paragraph(), " ".join(block))

    document.save(target)
    return target


if __name__ == "__main__":
    path = convert()
    print("written:", path)
